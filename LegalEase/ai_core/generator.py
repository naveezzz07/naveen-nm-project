from io import BytesIO
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from fpdf import FPDF

def sanitize_text(text: str) -> str:
    text = text.replace('"', '"').replace('"', '"')
    text = text.replace('‘', "'").replace('’', "'")
    return text

def format_docx(text: str, doc_type: str) -> BytesIO:
    doc = Document()
    try:
        doc.add_picture('Image/Logo.png', width=Inches(1.5))
        last_paragraph = doc.paragraphs[-1]
        last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    except Exception:
        pass
        
    heading = doc.add_heading(doc_type, 0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    for line in text.split('\n'):
        if line.strip():
            doc.add_paragraph(line.strip())
            
    try:
        section = doc.sections[0]
        footer = section.footer
        footer_para = footer.paragraphs[0]
        footer_para.text = "LegalEase Inc. | contact@legalease.com | All Rights Reserved."
        footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    except Exception:
        pass
            
    f = BytesIO()
    doc.save(f)
    f.seek(0)
    return f

class PDF(FPDF):
    def header(self):
        try:
            self.image('Image/Logo.png', x=95, y=8, w=20)
        except Exception:
            pass
        self.set_font('Arial', 'B', 15)
        self.cell(0, 30, 'LegalEase', 0, 1, 'C')
        self.ln(10)
        
    def footer(self):
        self.set_y(-15)
        self.set_font('Arial', 'I', 8)
        self.cell(0, 10, 'LegalEase Inc. | contact@legalease.com | All Rights Reserved.', 0, 0, 'C')

def format_pdf(text: str, doc_type: str) -> BytesIO:
    pdf = PDF()
    pdf.add_page()
    pdf.set_font('Arial', 'B', 16)
    
    safe_doc_type = doc_type.encode('latin-1', 'replace').decode('latin-1')
    pdf.cell(0, 10, safe_doc_type, 0, 1, 'C')
    pdf.ln(10)
    
    pdf.set_font('Arial', '', 12)
    safe_text = text.encode('latin-1', 'replace').decode('latin-1')
    
    for line in safe_text.split('\n'):
        if line.strip():
            pdf.multi_cell(0, 8, line.strip())
            pdf.ln(2)
            
    f = BytesIO()
    out = pdf.output(dest='S')
    if isinstance(out, str):
        f.write(out.encode('latin-1', 'replace'))
    else:
        f.write(out)
    f.seek(0)
    return f

def format_html_preview(text: str) -> str:
    html = "<div style='padding: 20px; background-color: #1e1e2e; color: #cdd6f4; border-radius: 10px;'>"
    for line in text.split('\n'):
        if line.strip().startswith('##'):
            html += f"<h3>{line.replace('##', '').strip()}</h3>"
        elif line.strip():
            html += f"<p>{line.strip()}</p>"
    html += "</div>"
    return html
